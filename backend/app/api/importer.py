"""Excel import: selected columns only + admin review of duplicate/invalid rows."""
import logging
import os
import re
import tempfile
from datetime import datetime, timezone
from uuid import UUID

from sqlalchemy.exc import IntegrityError

import openpyxl
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from pydantic import BaseModel
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.deps import admin_only
from app.db.session import get_db
from app.models import (
    ImportBatch, ImportError, Lead, LeadAssignment, LeadSource, LeadStatus, LeadStatusHistory,
    Product, ProductAlias, User,
)
from app.services import email_service, live
from app.services.lead_service import auto_assign, format_enquiry_number, next_enquiry_number
from app.services.pricing import apply_pricing_to_lead

log = logging.getLogger(__name__)
from app.services.normalize import (
    PRODUCT_ALIASES, SOURCE_ALIASES, canonical_source, is_valid_email, is_valid_phone,
    norm_key, norm_phone, parse_excel_date, parse_quantity,
)

router = APIRouter(prefix="/api/import", tags=["import"])
PENDING: dict[str, list[dict]] = {}

# Tracker sheet detection (order-free).
TRACKER_MARKERS = {"enq no", "date received", "lead name", "city", "lead source"}

# Excel Leads Tracker column indexes (0-based) — fallback when header lookup fails.
COL_ENQ, COL_DATE, COL_NAME, COL_COMPANY = 0, 1, 2, 3
COL_PHONE, COL_CITY, COL_CARS, COL_SOURCE, COL_PRODUCT = 4, 5, 6, 16, 19


def _resolve_cols(header: list[str]) -> dict[str, int]:
    """Map required fields to column indexes by header name, falling back to tracker defaults.

    Ensures the lead source (and other fields) are always taken from the
    uploaded Excel's own columns even if column order shifts.
    """
    lowered = [(h or "").strip().lower() for h in header]
    def find(keywords: list[str], fallback: int) -> int:
        for i, h in enumerate(lowered):
            if not h:
                continue
            for kw in keywords:
                if kw in h:
                    return i
        return fallback
    return {
        "enq": find(["enq"], COL_ENQ),
        "date": find(["date received", "received date", "date"], COL_DATE),
        "name": find(["lead name", "full name", "lead / full", "customer name", "name"], COL_NAME),
        "company": find(["company", "organisation", "organization"], COL_COMPANY),
        "phone": find(["contact no", "contact number", "phone", "mobile"], COL_PHONE),
        "requirement": find(["requirement"], 6),
        "quantity": find(["quantity"], 7),
        "remarks": find(["remarks", "remark"], 8),
        "priority": find(["priority"], 18),
        "email": find(["email id", "e-mail", "e mail", "email"], 29),
        "alternate_contact": find(["alternate", "alt contact"], 30),
        "city": find(["city"], COL_CITY),
        "cars": find(["no. of cars", "no of cars", "cars"], COL_CARS),
        "source": find(["lead source", "source"], COL_SOURCE),
        "product": find(["product", "type"], COL_PRODUCT),
    }


class ErrorUpdate(BaseModel):
    name: str | None = None
    city: str | None = None
    company: str | None = None
    phone: str | None = None
    cars: str | None = None
    requirement: str | None = None
    quantity: str | None = None
    remarks: str | None = None
    priority: str | None = None
    email: str | None = None
    alternate_contact: str | None = None
    source: str | None = None
    product: str | None = None
    enq: str | int | float | None = None
    date: str | None = None


class PromoteIn(BaseModel):
    force: bool = False


def parse_legacy_enq(raw) -> int | None:
    """Enquiry number from Excel or Google Sheets: 1, 1.0, '225', or 'ENQ-000225'."""
    if raw is None or raw == "":
        return None
    text = str(raw).strip()
    if not text:
        return None
    try:
        number = int(float(text))
    except (ValueError, TypeError):
        digits = re.sub(r"\D", "", text)
        if not digits:
            return None
        number = int(digits)
    if number < 1:
        return None
    return number


def suggest_sheet(names: list[str]) -> str:
    for n in names:
        k = n.lower()
        if "tracker" in k or ("lead" in k and "dashboard" not in k and "meta" not in k):
            return n
    return names[0] if names else ""


def header_score(header: list[str]) -> int:
    normed = {re.sub(r"\s+", " ", (h or "").strip().lower()) for h in header}
    hits = 0
    for m in TRACKER_MARKERS:
        if any(m in h or h in m for h in normed if h):
            hits += 1
    return hits


def _default_status(db: Session) -> LeadStatus:
    st = db.query(LeadStatus).filter(LeadStatus.name == "New Lead").first()
    if not st:
        raise HTTPException(500, "New Lead status missing — run seed")
    return st


def _norm_source(db: Session, raw: str) -> LeadSource | None:
    canon = canonical_source(raw)
    src = db.query(LeadSource).filter(func.lower(LeadSource.name) == canon.lower()).first()
    if src:
        return src
    return db.query(LeadSource).filter(func.lower(LeadSource.name) == "others").first()


def _norm_product(db: Session, raw: str) -> Product | None:
    if not raw or not str(raw).strip():
        return None
    k = norm_key(str(raw))
    canon = PRODUCT_ALIASES.get(k, str(raw).strip())
    p = db.query(Product).filter(
        func.lower(Product.name) == canon.lower(), Product.is_active.is_(True),
    ).first()
    if p:
        return p
    al = db.query(ProductAlias).filter(func.lower(ProductAlias.alias) == k).first()
    if not al:
        return None
    prod = db.get(Product, al.product_id)
    return prod if prod and prod.is_active else None


def _cell(r, idx):
    return r[idx] if len(r) > idx else None


def _json_safe_rec(rec: dict) -> dict:
    out = dict(rec)
    if "date" in out and out["date"] is not None and not isinstance(out["date"], (str, int, float, bool)):
        out["date"] = str(out["date"])
    out.pop("dup", None)
    return out


def _serialize_error(e: ImportError) -> dict:
    raw = e.raw or {}
    return {
        "id": str(e.id),
        "batch_id": str(e.batch_id),
        "row_number": e.row_number,
        "reason": e.reason,
        "error": e.error,
        "name": raw.get("name", ""),
        "city": raw.get("city", ""),
        "company": raw.get("company", ""),
        "phone": raw.get("phone", ""),
        "cars": raw.get("cars", ""),
        "requirement": raw.get("requirement", ""),
        "quantity": raw.get("quantity", ""),
        "remarks": raw.get("remarks", ""),
        "priority": raw.get("priority", ""),
        "email": raw.get("email", ""),
        "email_invalid": bool(raw.get("email_invalid", False)),
        "alternate_contact": raw.get("alternate_contact", ""),
        "source": raw.get("source", ""),
        "product": raw.get("product", ""),
        "enq": raw.get("legacy_enq", raw.get("enq")),
        "date": raw.get("date"),
        "raw": raw,
    }


def classify_intake_row(
    db: Session,
    *,
    name: str,
    phone: str,
    enq,
    email: str,
    seen_phones: set[str] | None = None,
    seen_enqs: set[int] | None = None,
) -> dict:
    """Shared Excel + Google Sheets rules.

    A valid phone or a valid email is enough. Name, enquiry number, and the
    other columns may be blank. A supplied enquiry number must still be unique.
    A duplicate phone is rejected. When the phone is valid, an unusable email
    is stored and does not block the row.
    """
    seen_phones = seen_phones or set()
    seen_enqs = seen_enqs or set()
    phone_n = norm_phone(phone or "")
    legacy = parse_legacy_enq(enq)
    email_v = (email or "").strip()
    phone_present = bool(str(phone or "").strip())
    phone_ok = is_valid_phone(phone or "")
    email_ok = bool(email_v) and is_valid_email(email_v)
    email_bad = bool(email_v) and not email_ok
    errs: list[str] = []
    dups: list[str] = []
    if phone_present and not phone_ok:
        shown = phone_n or str(phone or "").strip()
        errs.append(f"missing/invalid phone ({shown})")
    elif phone_ok and (phone_n in seen_phones or db.query(Lead).filter_by(contact_number_norm=phone_n).first()):
        dups.append("duplicate phone")
    elif not phone_ok and not email_ok:
        if email_bad:
            errs.append("invalid email")
        else:
            errs.append("missing phone or email")
    if legacy is not None and (
        legacy in seen_enqs
        or db.query(Lead).filter_by(legacy_enquiry_no=legacy).first()
        or db.query(Lead).filter_by(enquiry_number=format_enquiry_number(legacy)).first()
    ):
        dups.append(f"duplicate enquiry no {format_enquiry_number(legacy)}")
    return {
        "phone_norm": phone_n if phone_ok else "",
        "legacy_enq": legacy,
        "email_invalid": email_bad,
        "dups": dups,
        "errs": errs,
        "reason": "DUPLICATE" if dups else ("INVALID" if errs else "OK"),
        "messages": dups or errs,
    }


def _row_issue_messages(db: Session, raw: dict) -> list[str]:
    classified = classify_intake_row(
        db,
        name=str(raw.get("name") or ""),
        phone=str(raw.get("phone") or ""),
        enq=raw.get("enq") if raw.get("legacy_enq") is None else raw.get("legacy_enq"),
        email=str(raw.get("email") or ""),
    )
    return classified["messages"]


def _assigned_date_str(db: Session, lead: Lead) -> str:
    """Assigned date for emails: current assignment row, else lead creation."""
    try:
        assign = db.query(LeadAssignment).filter(
            LeadAssignment.lead_id == lead.id,
            LeadAssignment.is_current.is_(True),
        ).order_by(LeadAssignment.assigned_at.desc()).first()
        at = assign.assigned_at if assign else lead.created_at
        return at.strftime("%d-%b-%Y") if at else "—"
    except Exception:
        return "—"


def _lead_email_item_db(lead: Lead, src_map: dict, prod_map: dict, assign_map: dict) -> dict:
    """Bulk-prefetched variant of _lead_email_item (no per-lead queries)."""
    assign = assign_map.get(lead.id)
    at = assign.assigned_at if assign else lead.created_at
    deadline = lead.sla_deadline.strftime("%d-%b-%Y %H:%M") if lead.sla_deadline else "—"
    return {
        "enquiry_number": lead.enquiry_number,
        "legacy_enq": lead.legacy_enquiry_no,
        "enquiry_date": str(lead.enquiry_date) if lead.enquiry_date else "—",
        "assigned_date_str": at.strftime("%d-%b-%Y") if at else "—",
        "customer_name": lead.customer_name or "",
        "contact_number": lead.contact_number or "",
        "alternate_contact": lead.alternate_contact or "",
        "email": lead.email or "",
        "company_name": lead.company_name or "",
        "city": lead.city or "",
        "source": src_map.get(lead.source_id, "—") if lead.source_id else "—",
        "product": (lead.product_raw or prod_map.get(lead.product_id) or "—"),
        "quantity_raw": lead.quantity_raw or "",
        "deadline_str": deadline,
        "lead_url": f"{settings.FRONTEND_URL.rstrip('/')}/leads/{lead.id}",
    }


def _lead_email_item(db: Session, lead: Lead) -> dict:
    """Full customer details payload for the batched assignment email."""
    src = db.get(LeadSource, lead.source_id) if lead.source_id else None
    prod = db.get(Product, lead.product_id) if lead.product_id else None
    deadline = lead.sla_deadline.strftime("%d-%b-%Y %H:%M") if lead.sla_deadline else "—"
    return {
        "enquiry_number": lead.enquiry_number,
        "legacy_enq": lead.legacy_enquiry_no,
        "enquiry_date": str(lead.enquiry_date) if lead.enquiry_date else "—",
        "assigned_date_str": _assigned_date_str(db, lead),
        "customer_name": lead.customer_name or "",
        "contact_number": lead.contact_number or "",
        "alternate_contact": lead.alternate_contact or "",
        "email": lead.email or "",
        "company_name": lead.company_name or "",
        "city": lead.city or "",
        "source": src.name if src else "—",
        "product": (lead.product_raw or (prod.name if prod else "") or "—"),
        "quantity_raw": lead.quantity_raw or "",
        "deadline_str": deadline,
        "lead_url": f"{settings.FRONTEND_URL.rstrip('/')}/leads/{lead.id}",
    }


def _send_assignment_batches(db: Session, new_by_emp: dict) -> None:
    """Send one batched assignment email per employee. Fail-open (never raises)."""
    now = datetime.now(timezone.utc)
    try:
        all_ids = [lid for ids in (new_by_emp or {}).values() for lid in ids]
        leads_by_id = {x.id: x for x in db.query(Lead).filter(Lead.id.in_(all_ids)).all()} if all_ids else {}
        src_map = {s.id: s.name for s in db.query(LeadSource).all()}
        prod_map = {p.id: p.name for p in db.query(Product).all()}
        assigns = db.query(LeadAssignment).filter(
            LeadAssignment.lead_id.in_(list(leads_by_id)),
            LeadAssignment.is_current.is_(True),
        ).all() if leads_by_id else []
        assign_map = {a.lead_id: a for a in assigns}
        for emp_id, lead_ids in (new_by_emp or {}).items():
            try:
                emp = db.get(User, emp_id)
                if not emp:
                    continue
                leads = [leads_by_id[lid] for lid in lead_ids if lid in leads_by_id]
                if not leads:
                    continue
                items = [_lead_email_item_db(x, src_map, prod_map, assign_map) for x in leads]
                sent = email_service.send_assignment_email(emp.email, emp.name, items)
                if sent:
                    for x in leads:
                        x.assignment_email_sent_at = now
                    db.commit()
            except Exception as exc:
                log.error("assignment email batch failed for %s: %s", emp_id, exc)
                try:
                    db.rollback()
                except Exception:
                    pass
    except Exception as exc:
        log.error("assignment email dispatch failed: %s", exc)


def _create_lead_from_raw(db: Session, raw: dict, admin: User, *, force: bool = False) -> Lead:
    name = str(raw.get("name") or "").strip()
    phone = str(raw.get("phone") or "").strip()
    email_v = str(raw.get("email") or "").strip()
    phone_ok = is_valid_phone(phone)
    email_ok = bool(email_v) and is_valid_email(email_v)
    if phone and not phone_ok:
        raise HTTPException(400, "Phone must be a 10-digit Indian number")
    if not phone_ok and not email_ok:
        raise HTTPException(400, "A valid phone or email is required")
    phone_n = norm_phone(phone) if phone_ok else ""
    if phone_ok and db.query(Lead).filter_by(contact_number_norm=phone_n).first():
        raise HTTPException(400, "Phone already exists on another lead — correct the phone first")

    legacy = raw.get("legacy_enq")
    if legacy is None:
        legacy = parse_legacy_enq(raw.get("enq"))
    if legacy is None:
        enquiry_number = next_enquiry_number(db)
        legacy = None
    else:
        enquiry_number = format_enquiry_number(int(legacy))
        if (
            db.query(Lead).filter_by(legacy_enquiry_no=legacy).first()
            or db.query(Lead).filter_by(enquiry_number=enquiry_number).first()
        ):
            raise HTTPException(400, f"Enquiry number {enquiry_number} already exists")

    company = str(raw.get("company") or "").strip()
    city = str(raw.get("city") or "").strip()
    cars = str(raw.get("cars") or "").strip()
    src = _norm_source(db, str(raw.get("source") or ""))
    prod = _norm_product(db, str(raw.get("product") or ""))
    st = _default_status(db)

    try:
        with db.begin_nested():
            lead = Lead(
                enquiry_number=enquiry_number,
                legacy_enquiry_no=legacy,
                enquiry_date=parse_excel_date(raw.get("date")),
                customer_name=name,
                company_name=company,
                contact_number=phone,
                contact_number_norm=phone_n,
                alternate_contact=str(raw.get("alternate_contact") or "").strip(),
                email=str(raw.get("email") or "").strip(),
                city=city,
                product_raw=str(raw.get("product") or "").strip(),
                requirement=str(raw.get("requirement") or "").strip(),
                quantity_raw=cars or str(raw.get("quantity") or "").strip(),
                quantity_num=parse_quantity(cars or raw.get("quantity")),
                priority=str(raw.get("priority") or "").strip(),
                first_contact_notes=str(raw.get("remarks") or "").strip(),
                source_id=src.id if src else None,
                product_id=prod.id if prod else None,
                status_id=st.id,
                sla_state="PENDING",
                created_by=admin.id,
            )
            apply_pricing_to_lead(lead, product=prod)
            db.add(lead)
            db.flush()
            db.add(LeadStatusHistory(
                lead_id=lead.id, old_status_id=None, new_status_id=st.id,
                changed_by=admin.id, reason="import review promote",
            ))
            auto_assign(db, lead, admin)
    except IntegrityError:
        raise HTTPException(400, "Enquiry number already exists")
    return lead


@router.post("/excel")
async def upload_excel(file: UploadFile = File(...), sheet: str = Form(""),
                       db: Session = Depends(get_db), admin: User = Depends(admin_only)):
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(400, "Please upload an .xlsx file")
    data = await file.read()
    if not data:
        raise HTTPException(400, "Uploaded file is empty")
    safe_name = re.sub(r"[^A-Za-z0-9_.-]", "_", os.path.basename(file.filename))
    tmp = tempfile.NamedTemporaryFile(prefix="crm_import_", suffix=".xlsx", delete=False)
    try:
        tmp.write(data)
        tmp.close()
        wb = openpyxl.load_workbook(tmp.name, data_only=True, read_only=True)
    finally:
        try:
            os.unlink(tmp.name)
        except OSError:
            pass
    names = wb.sheetnames
    if sheet and sheet not in names:
        raise HTTPException(400, f"Sheet '{sheet}' not found. Available: {', '.join(names)}")
    if not sheet:
        return {"batch_id": None, "sheets": names, "suggested": suggest_sheet(names)}
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    if not rows or not any(rows[0]):
        raise HTTPException(400, f"Sheet '{sheet}' is empty")
    header = [str(c or "").strip() for c in rows[0]]
    if header_score(header) < 3:
        raise HTTPException(
            400,
            f"Sheet '{sheet}' does not look like the Leads Tracker "
            f"(found headers: {', '.join(h for h in header if h)[:120]}). "
            "Please select the tracker sheet.")

    preview, duplicates, invalids = [], [], []
    seen_enqs: set[int] = set()
    seen_phones: set[str] = set()
    cols = _resolve_cols(header)
    for i, r in enumerate(rows[1:], start=2):
        if not any(r):
            continue
        date_v = _cell(r, cols["date"])
        if date_v is not None and not isinstance(date_v, (str, int, float, bool)):
            date_v = str(date_v)
        rec = {
            "row": i,
            "enq": _cell(r, cols["enq"]),
            "date": date_v,
            "name": str(_cell(r, cols["name"]) or "").strip(),
            "company": str(_cell(r, cols["company"]) or "").strip(),
            "phone": str(_cell(r, cols["phone"]) or "").strip(),
            "city": str(_cell(r, cols["city"]) or "").strip(),
            "cars": str(_cell(r, cols["cars"]) or "").strip(),
            "requirement": str(_cell(r, cols["requirement"]) or "").strip(),
            "quantity": str(_cell(r, cols["quantity"]) or "").strip(),
            "remarks": str(_cell(r, cols["remarks"]) or "").strip(),
            "priority": str(_cell(r, cols["priority"]) or "").strip(),
            "email": str(_cell(r, cols["email"]) or "").strip(),
            "alternate_contact": str(_cell(r, cols["alternate_contact"]) or "").strip(),
            "source": str(_cell(r, cols["source"]) or "").strip(),
            "product": str(_cell(r, cols["product"]) or "").strip(),
        }
        rec["legacy_enq"] = None
        rec["phone_norm"] = ""
        rec["product_unmapped"] = bool((rec.get("product") or "").strip()) and _norm_product(db, rec["product"]) is None
        classified = classify_intake_row(
            db, name=rec["name"], phone=rec["phone"], enq=rec["enq"], email=rec["email"],
            seen_phones=seen_phones, seen_enqs=seen_enqs,
        )
        rec["legacy_enq"] = classified["legacy_enq"]
        rec["phone_norm"] = classified["phone_norm"]
        rec["email_invalid"] = classified["email_invalid"]

        if classified["dups"]:
            rec["dup"] = True
            rec["errors"] = classified["dups"]
            duplicates.append(rec)
        elif classified["errs"]:
            rec["errors"] = classified["errs"]
            invalids.append(rec)
        else:
            if classified["phone_norm"]:
                seen_phones.add(classified["phone_norm"])
            if classified["legacy_enq"] is not None:
                seen_enqs.add(classified["legacy_enq"])
        preview.append(rec)

    batch = ImportBatch(
        file_name=safe_name, sheet_name=ws.title, total_rows=len(preview),
        duplicates=len(duplicates), invalid=len(invalids), status="PREVIEW", created_by=admin.id,
    )
    db.add(batch)
    db.commit()
    PENDING[str(batch.id)] = preview
    valid = len(preview) - len(duplicates) - len(invalids)
    return {
        "batch_id": str(batch.id), "sheets": names, "sheet": ws.title,
        "total": len(preview), "duplicates": len(duplicates), "invalid": len(invalids),
        "valid": valid,
        "preview": [r for r in preview if not r.get("dup") and not r.get("errors")][:50],
        "duplicate_rows": duplicates,
        "invalid_rows": invalids,
        "fields": [
            "enquiry no", "received date", "name", "company/organisation (optional)",
            "contact no", "city", "no. of cars", "lead source", "product/type", "email",
        ],
        "note": "A valid phone or a valid email is enough. Other columns may be blank. A duplicate phone, or a duplicate enquiry number when one is filled in, is held for review. Unmapped product names still import as-is.",
    }


@router.post("/{bid}/confirm")
def confirm(bid: UUID, db: Session = Depends(get_db), admin: User = Depends(admin_only)):
    batch = db.get(ImportBatch, bid)
    rows = PENDING.get(str(bid), [])
    if not batch or not rows:
        raise HTTPException(404, "Batch expired — re-upload")
    st = _default_status(db)
    ok = assigned = pending = 0
    new_by_emp: dict = {}
    for rec in rows:
        if rec.get("dup") or rec.get("errors"):
            db.add(ImportError(
                batch_id=bid, row_number=rec["row"], raw=_json_safe_rec(rec),
                error=",".join(rec.get("errors", ["duplicate"])),
                reason="DUPLICATE" if rec.get("dup") else "INVALID",
            ))
            continue
        legacy = rec.get("legacy_enq")
        if legacy is None:
            enquiry_number = next_enquiry_number(db)
        else:
            enquiry_number = format_enquiry_number(int(legacy))
            if (
                db.query(Lead).filter_by(legacy_enquiry_no=legacy).first()
                or db.query(Lead).filter_by(enquiry_number=enquiry_number).first()
            ):
                db.add(ImportError(
                    batch_id=bid, row_number=rec["row"], raw=_json_safe_rec(rec),
                    error=f"duplicate enquiry no {enquiry_number}", reason="DUPLICATE",
                ))
                batch.duplicates += 1
                continue
        phone_n = rec.get("phone_norm") or norm_phone(str(rec.get("phone") or ""))
        if phone_n and db.query(Lead).filter_by(contact_number_norm=phone_n).first():
            db.add(ImportError(
                batch_id=bid, row_number=rec["row"], raw=_json_safe_rec(rec),
                error="duplicate phone", reason="DUPLICATE",
            ))
            batch.duplicates += 1
            continue

        cars = str(rec.get("cars") or "")
        phone = str(rec.get("phone") or "")
        phone_n = rec.get("phone_norm") or norm_phone(phone)
        src = _norm_source(db, str(rec.get("source") or ""))
        prod = _norm_product(db, str(rec.get("product") or ""))
        try:
            with db.begin_nested():
                lead = Lead(
                    enquiry_number=enquiry_number,
                    legacy_enquiry_no=legacy,
                    enquiry_date=parse_excel_date(rec.get("date")),
                    customer_name=str(rec.get("name") or ""),
                    company_name=str(rec.get("company") or ""),
                    contact_number=phone,
                    contact_number_norm=phone_n,
                    alternate_contact=str(rec.get("alternate_contact") or "").strip(),
                    email=str(rec.get("email") or "").strip(),
                    city=str(rec.get("city") or ""),
                    product_raw=str(rec.get("product") or "").strip(),
                    requirement=str(rec.get("requirement") or "").strip(),
                    quantity_raw=cars or str(rec.get("quantity") or "").strip(),
                    quantity_num=parse_quantity(cars or rec.get("quantity")),
                    priority=str(rec.get("priority") or "").strip(),
                    first_contact_notes=str(rec.get("remarks") or "").strip(),
                    source_id=src.id if src else None,
                    product_id=prod.id if prod else None,
                    status_id=st.id,
                    sla_state="PENDING",
                    created_by=admin.id,
                )
                apply_pricing_to_lead(lead, product=prod)
                db.add(lead)
                db.flush()
                db.add(LeadStatusHistory(
                    lead_id=lead.id, old_status_id=None, new_status_id=st.id,
                    changed_by=admin.id, reason="excel import",
                ))
                emp = auto_assign(db, lead, admin)
                if emp:
                    new_by_emp.setdefault(emp.id, []).append(lead.id)
        except IntegrityError:
            db.add(ImportError(
                batch_id=bid, row_number=rec["row"], raw=_json_safe_rec(rec),
                error="duplicate enquiry no or phone", reason="DUPLICATE",
            ))
            batch.duplicates += 1
            continue
        if emp:
            assigned += 1
        else:
            pending += 1
        ok += 1
    batch.imported = ok
    batch.status = "DONE"
    db.commit()
    if ok:
        live.bump()
    # One batched email per employee with all newly assigned customers.
    _send_assignment_batches(db, new_by_emp)
    errors = db.query(ImportError).filter_by(batch_id=bid).order_by(ImportError.row_number).all()
    return {
        "batch_id": str(bid),
        "total": batch.total_rows,
        "imported": ok,
        "assigned": assigned,
        "pending": pending,
        "duplicates": batch.duplicates,
        "invalid": batch.invalid,
        "skipped": len(errors),
        "errors": [_serialize_error(e) for e in errors],
    }


@router.get("/batches")
def list_batches(db: Session = Depends(get_db), _: User = Depends(admin_only), limit: int = 20):
    rows = (
        db.query(ImportBatch)
        .filter(ImportBatch.status != "SYNCED")
        .order_by(ImportBatch.created_at.desc())
        .limit(limit)
        .all()
    )
    return [{
        "id": str(b.id), "file_name": b.file_name, "sheet_name": b.sheet_name,
        "status": b.status, "total_rows": b.total_rows, "imported": b.imported,
        "duplicates": b.duplicates, "invalid": b.invalid,
        "source": "sheets" if (b.file_name or "").startswith("sheets:") else "excel",
        "created_at": b.created_at.isoformat() if b.created_at else None,
    } for b in rows]


@router.get("/{bid}/errors")
def list_errors(bid: UUID, reason: str = "", db: Session = Depends(get_db), _: User = Depends(admin_only)):
    q = db.query(ImportError).filter_by(batch_id=bid)
    if reason:
        q = q.filter(ImportError.reason == reason.upper())
    rows = q.order_by(ImportError.row_number).all()
    return {
        "batch_id": str(bid),
        "total": len(rows),
        "duplicates": sum(1 for e in rows if e.reason == "DUPLICATE"),
        "invalid": sum(1 for e in rows if e.reason == "INVALID"),
        "items": [_serialize_error(e) for e in rows],
    }


@router.patch("/errors/{eid}")
def update_error(eid: UUID, body: ErrorUpdate, db: Session = Depends(get_db), _: User = Depends(admin_only)):
    e = db.get(ImportError, eid)
    if not e:
        raise HTTPException(404, "Skipped row not found")
    raw = dict(e.raw or {})
    data = body.model_dump(exclude_unset=True)
    for key in ("name", "city", "company", "phone", "cars", "requirement", "quantity",
                "remarks", "priority", "email", "alternate_contact", "source", "product", "date"):
        if key in data and data[key] is not None:
            raw[key] = str(data[key]).strip() if key != "date" else data[key]
    if "phone" in data and data["phone"] is not None:
        raw["phone_norm"] = norm_phone(str(data["phone"]))
    if "email" in data and data["email"] is not None:
        email_v = str(data["email"]).strip()
        raw["email_invalid"] = bool(email_v) and not is_valid_email(email_v)
    if "enq" in data and data["enq"] is not None:
        raw["enq"] = data["enq"]
        raw["legacy_enq"] = parse_legacy_enq(data["enq"])
    e.raw = raw
    msgs = _row_issue_messages(db, raw)
    e.error = ", ".join(msgs) if msgs else "ready to add"
    if not msgs:
        e.reason = "READY"
    else:
        e.reason = "DUPLICATE" if any("duplicate" in m for m in msgs) else "INVALID"
    db.commit()
    db.refresh(e)
    return _serialize_error(e)


@router.post("/errors/{eid}/promote")
def promote_error(eid: UUID, body: PromoteIn | None = None, db: Session = Depends(get_db),
                  admin: User = Depends(admin_only)):
    e = db.get(ImportError, eid)
    if not e:
        raise HTTPException(404, "Skipped row not found")
    force = bool(body and body.force)
    lead = _create_lead_from_raw(db, e.raw or {}, admin, force=force)
    emp_id = lead.primary_employee_id
    lead_id = lead.id
    batch = db.get(ImportBatch, e.batch_id)
    if batch:
        batch.imported = (batch.imported or 0) + 1
        if e.reason == "DUPLICATE" and batch.duplicates:
            batch.duplicates = max(0, batch.duplicates - 1)
        if e.reason == "INVALID" and batch.invalid:
            batch.invalid = max(0, batch.invalid - 1)
    db.delete(e)
    db.commit()
    live.bump()
    if emp_id:
        _send_assignment_batches(db, {emp_id: [lead_id]})
    return {
        "ok": True,
        "lead_id": str(lead.id),
        "enquiry_number": lead.enquiry_number,
        "assigned": bool(lead.primary_employee_id),
    }


@router.delete("/errors/{eid}")
def dismiss_error(eid: UUID, db: Session = Depends(get_db), _: User = Depends(admin_only)):
    e = db.get(ImportError, eid)
    if not e:
        raise HTTPException(404, "Skipped row not found")
    db.delete(e)
    db.commit()
    return {"ok": True}
